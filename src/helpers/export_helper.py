# src/helpers/export_helper.py
import io
from typing import Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

def gerar_excel_relatorio(rel: Dict[str, Any]) -> bytes:
    """Gera planilha Excel (.xlsx) estilizada do relatório bimestral."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório de Gestão"

    # Estilos
    font_header = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=11, bold=True, color="1F2937")
    font_sec = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)

    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue_light = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    fill_green = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_gray = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color="D1D5DB"),
        right=Side(style='thin', color="D1D5DB"),
        top=Side(style='thin', color="D1D5DB"),
        bottom=Side(style='thin', color="D1D5DB")
    )

    # 1. Título e Identificação
    ws.merge_cells("A1:G1")
    ws["A1"] = "HOSPITAL DAS CLÍNICAS DA UFPE (HC-UFPE / EBSERH)"
    ws["A1"].font = font_header
    ws["A1"].fill = fill_navy
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Relatório de Gestão — {rel.get('bimestre_nome', 'Bimestre')} / Exercício {rel.get('exercicio_id', 2026)}"
    ws["A2"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    ws["A2"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Metadados
    row = 4
    info_pairs = [
        ("Unidade Organizacional (UORG):", f"{rel.get('uorg_sigla', '')} - {rel.get('uorg_nome', '')}"),
        ("Status do Relatório:", rel.get("status", "RASCUNHO")),
        ("Data de Submissão:", rel.get("data_envio", "Não submetido")),
        ("Responsável pelo Envio:", rel.get("usuario_envio", "—"))
    ]
    for label, val in info_pairs:
        ws.cell(row=row, column=1, value=label).font = font_sub
        ws.cell(row=row, column=2, value=val).font = font_regular
        row += 1

    # 2. Entregas Planejadas
    row += 1
    ws.merge_cells(f"A{row}:G{row}")
    ws[f"A{row}"] = "1. ENTREGAS PLANEJADAS PACTUADAS"
    ws[f"A{row}"].font = font_sec
    row += 1

    headers_ent = ["#", "Título da Entrega", "Concluída no Bimestre?", "Observações"]
    ws.cell(row=row, column=1, value=headers_ent[0]).fill = fill_blue_light
    ws.cell(row=row, column=2, value=headers_ent[1]).fill = fill_blue_light
    ws.cell(row=row, column=3, value=headers_ent[2]).fill = fill_blue_light
    ws.cell(row=row, column=4, value=headers_ent[3]).fill = fill_blue_light
    for col_idx in range(1, 5):
        ws.cell(row=row, column=col_idx).font = font_bold
        ws.cell(row=row, column=col_idx).border = thin_border
    row += 1

    itens_ent = rel.get("itens_entrega", [])
    if not itens_ent:
        ws.cell(row=row, column=2, value="Nenhuma entrega pendente cadastrada.").font = font_regular
        row += 1
    else:
        for idx, ent in enumerate(itens_ent, 1):
            conc_str = "SIM" if ent.get("concluida") else "NÃO"
            ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2, value=ent.get("titulo", ""))
            ws.cell(row=row, column=3, value=conc_str).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=4, value=ent.get("observacoes", ""))

            fill_status = fill_green if ent.get("concluida") else fill_gray
            ws.cell(row=row, column=3).fill = fill_status

            for c in range(1, 5):
                ws.cell(row=row, column=c).font = font_regular
                ws.cell(row=row, column=c).border = thin_border
            row += 1

    # 3. Indicadores de Gestão
    row += 1
    ws.merge_cells(f"A{row}:G{row}")
    ws[f"A{row}"] = "2. INDICADORES DE GESTÃO"
    ws[f"A{row}"].font = font_sec
    row += 1

    headers_ind = ["#", "Nome do Indicador", "Periodicidade", "Meta", "Valor Apurado", "Status da Meta", "Justificativa"]
    for idx, h in enumerate(headers_ind, 1):
        cell = ws.cell(row=row, column=idx, value=h)
        cell.font = font_bold
        cell.fill = fill_blue_light
        cell.border = thin_border
    row += 1

    itens_ind = rel.get("itens_indicador", [])
    for idx, ind in enumerate(itens_ind, 1):
        op = ind.get("meta_operador", ">=")
        meta_txt = f"{op} {ind.get('meta_valor', '')}"
        if op == "ENTRE":
            meta_txt = f"{ind.get('meta_valor', '')} a {ind.get('meta_valor_max', '')}"

        if ind.get("resultado_indisponivel"):
            apurado_txt = "Indisponível"
            status_txt = "Justificado"
            fill_ind = fill_gray
        elif ind.get("nao_aplicavel"):
            apurado_txt = "N/A"
            status_txt = "Não Aplicável"
            fill_ind = fill_gray
        else:
            apurado_txt = str(ind.get("valor_apurado", ""))
            atingida = ind.get("meta_atingida")
            if atingida is True:
                status_txt = "Atingida"
                fill_ind = fill_green
            elif atingida is False:
                status_txt = "Não Atingida"
                fill_ind = fill_red
            else:
                status_txt = "Pendente"
                fill_ind = fill_gray

        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=ind.get("nome", ""))
        ws.cell(row=row, column=3, value=ind.get("periodicidade", "")).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=4, value=meta_txt).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=5, value=apurado_txt).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=6, value=status_txt).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=6).fill = fill_ind
        ws.cell(row=row, column=7, value=ind.get("justificativa_indisponibilidade", ""))

        for c in range(1, 8):
            ws.cell(row=row, column=c).font = font_regular
            ws.cell(row=row, column=c).border = thin_border
        row += 1

    # 4. Destaques e Dificuldades
    row += 1
    ws.merge_cells(f"A{row}:G{row}")
    ws[f"A{row}"] = "3. DESTAQUES OPERACIONAIS"
    ws[f"A{row}"].font = font_sec
    row += 1
    ws.merge_cells(f"A{row}:G{row+1}")
    ws[f"A{row}"] = rel.get("destaques") or "(Nenhum destaque informado)"
    ws[f"A{row}"].font = font_regular
    ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    row += 3

    ws.merge_cells(f"A{row}:G{row}")
    ws[f"A{row}"] = "4. PRINCIPAIS DIFICULDADES E ENTRAVES"
    ws[f"A{row}"].font = font_sec
    row += 1
    ws.merge_cells(f"A{row}:G{row+1}")
    ws[f"A{row}"] = rel.get("dificuldades") or "(Nenhuma dificuldade informada)"
    ws[f"A{row}"].font = font_regular
    ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    row += 3

    # 5. Feedback Institucional
    feedbacks = rel.get("feedbacks", [])
    if feedbacks:
        ws.merge_cells(f"A{row}:G{row}")
        ws[f"A{row}"] = "5. PARECER / FEEDBACK INSTITUCIONAL"
        ws[f"A{row}"].font = font_sec
        row += 1
        for fb in feedbacks:
            ws.merge_cells(f"A{row}:G{row}")
            ws[f"A{row}"] = f"{fb.get('autor_nome')} ({fb.get('autor_perfil')}) — {fb.get('data', '')}:"
            ws[f"A{row}"].font = font_bold
            row += 1
            ws.merge_cells(f"A{row}:G{row+1}")
            ws[f"A{row}"] = fb.get("texto", "")
            ws[f"A{row}"].font = font_regular
            ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
            row += 3

    # Ajustar largura de colunas
    col_widths = {1: 8, 2: 45, 3: 16, 4: 18, 5: 16, 6: 18, 7: 35}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[chr(64 + col_idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def gerar_excel_matriz(matriz: Dict[str, Any]) -> bytes:
    """Gera planilha Excel consolidada da Matriz UORG x Bimestre."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz de Acompanhamento"

    font_header = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)

    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_enviado = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_pendente = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_edicao = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    fill_prazo = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_fechado = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color="D1D5DB"),
        right=Side(style='thin', color="D1D5DB"),
        top=Side(style='thin', color="D1D5DB"),
        bottom=Side(style='thin', color="D1D5DB")
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = f"HC-UFPE — Painel Consolidado de Relatórios de Gestão (Exercício {matriz.get('exercicio_id', 2026)})"
    ws["A1"].font = font_header
    ws["A1"].fill = fill_navy
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    headers = ["Gerência", "Sigla UORG", "1º Bimestre", "2º Bimestre", "3º Bimestre", "4º Bimestre", "5º Bimestre", "6º Bimestre"]
    for idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=idx, value=h)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 22

    row = 4
    for linha in matriz.get("linhas", []):
        ws.cell(row=row, column=1, value=linha.get("gerencia_sigla", "")).font = font_bold
        ws.cell(row=row, column=2, value=linha.get("uorg_sigla", "")).font = font_bold
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=2).border = thin_border

        for col_idx, bim in enumerate(linha.get("bimestres", []), 3):
            status = bim.get("status", "NAO_INICIADO")
            cell = ws.cell(row=row, column=col_idx, value=status)
            cell.font = font_regular
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

            if status == "ENVIADO":
                cell.fill = fill_enviado
            elif status == "PENDENTE":
                cell.fill = fill_pendente
            elif status == "EM_EDICAO":
                cell.fill = fill_edicao
            elif status == "PRAZO_EXCEPCIONAL":
                cell.fill = fill_prazo
            else:
                cell.fill = fill_fechado

        row += 1

    col_widths = {1: 18, 2: 18, 3: 16, 4: 16, 5: 16, 6: 16, 7: 16, 8: 16}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[chr(64 + col_idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def gerar_pdf_relatorio(rel: Dict[str, Any]) -> bytes:
    """Gera documento PDF formatado institucionalmente do relatório bimestral."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1E3A8A"),
        alignment=1,
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Heading2'],
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#2563EB"),
        alignment=1,
        spaceAfter=12
    )
    sec_style = ParagraphStyle(
        'DocSec',
        parent=styles['Heading3'],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1F2937")
    )
    bold_style = ParagraphStyle(
        'DocBold',
        parent=body_style,
        fontName="Helvetica-Bold"
    )

    story = []

    # Cabeçalho
    story.append(Paragraph("HOSPITAL DAS CLÍNICAS DA UFPE (HC-UFPE / EBSERH)", title_style))
    story.append(Paragraph(f"Portal de Relatórios de Gestão — {rel.get('bimestre_nome', 'Bimestre')} / {rel.get('exercicio_id', 2026)}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1E3A8A"), spaceAfter=10))

    # Identificação
    meta_data = [
        [Paragraph("<b>Unidade (UORG):</b>", body_style), Paragraph(f"{rel.get('uorg_sigla', '')} - {rel.get('uorg_nome', '')}", body_style)],
        [Paragraph("<b>Status / Versão:</b>", body_style), Paragraph(f"{rel.get('status', 'RASCUNHO')} (v{rel.get('versao', 1)})", body_style)],
        [Paragraph("<b>Data de Submissão:</b>", body_style), Paragraph(rel.get("data_envio") or "Não submetido", body_style)],
        [Paragraph("<b>Responsável:</b>", body_style), Paragraph(rel.get("usuario_envio") or "—", body_style)]
    ]
    t_meta = Table(meta_data, colWidths=[120, 400])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 10))

    # 1. Entregas
    story.append(Paragraph("1. Entregas Planejadas Pactuadas", sec_style))
    itens_ent = rel.get("itens_entrega", [])
    if itens_ent:
        ent_rows = [[Paragraph("<b>Título da Entrega</b>", bold_style), Paragraph("<b>Concluída?</b>", bold_style), Paragraph("<b>Observações</b>", bold_style)]]
        for ent in itens_ent:
            c_txt = "<font color='green'><b>SIM</b></font>" if ent.get("concluida") else "<font color='gray'>NÃO</font>"
            ent_rows.append([
                Paragraph(ent.get("titulo", ""), body_style),
                Paragraph(c_txt, body_style),
                Paragraph(ent.get("observacoes", "") or "—", body_style)
            ])
        t_ent = Table(ent_rows, colWidths=[280, 80, 160])
        t_ent.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#E0F2FE")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_ent)
    else:
        story.append(Paragraph("Nenhuma entrega cadastrada.", body_style))

    story.append(Spacer(1, 10))

    # 2. Indicadores
    story.append(Paragraph("2. Indicadores de Gestão", sec_style))
    itens_ind = rel.get("itens_indicador", [])
    if itens_ind:
        ind_rows = [[
            Paragraph("<b>Indicador</b>", bold_style),
            Paragraph("<b>Meta</b>", bold_style),
            Paragraph("<b>Apurado</b>", bold_style),
            Paragraph("<b>Atingimento</b>", bold_style)
        ]]
        for ind in itens_ind:
            op = ind.get("meta_operador", ">=")
            meta_str = f"{op} {ind.get('meta_valor', '')}"
            if op == "ENTRE":
                meta_str = f"{ind.get('meta_valor', '')} a {ind.get('meta_valor_max', '')}"

            if ind.get("resultado_indisponivel"):
                ap_str = "Indisponível"
                st_str = "<font color='gray'>Justificado</font>"
            else:
                ap_str = str(ind.get("valor_apurado") if ind.get("valor_apurado") is not None else "—")
                at = ind.get("meta_atingida")
                if at is True:
                    st_str = "<font color='green'><b>Atingida</b></font>"
                elif at is False:
                    st_str = "<font color='red'><b>Não Atingida</b></font>"
                else:
                    st_str = "<font color='gray'>Pendente</font>"

            ind_rows.append([
                Paragraph(ind.get("nome", ""), body_style),
                Paragraph(meta_str, body_style),
                Paragraph(ap_str, body_style),
                Paragraph(st_str, body_style)
            ])
        t_ind = Table(ind_rows, colWidths=[240, 90, 90, 100])
        t_ind.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#E0F2FE")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_ind)

    story.append(Spacer(1, 10))

    # 3. Destaques e Dificuldades
    story.append(Paragraph("3. Destaques Operacionais", sec_style))
    story.append(Paragraph(rel.get("destaques") or "<i>Nenhum destaque informado.</i>", body_style))
    story.append(Spacer(1, 8))

    story.append(Paragraph("4. Principais Dificuldades e Entraves", sec_style))
    story.append(Paragraph(rel.get("dificuldades") or "<i>Nenhuma dificuldade informada.</i>", body_style))
    story.append(Spacer(1, 10))

    # 4. Feedbacks
    feedbacks = rel.get("feedbacks", [])
    if feedbacks:
        story.append(Paragraph("5. Parecer / Feedback Institucional", sec_style))
        for fb in feedbacks:
            story.append(Paragraph(f"<b>{fb.get('autor_nome')} ({fb.get('autor_perfil')}):</b>", body_style))
            story.append(Paragraph(fb.get("texto", ""), body_style))
            story.append(Spacer(1, 4))

    doc.build(story)
    return buf.getvalue()
